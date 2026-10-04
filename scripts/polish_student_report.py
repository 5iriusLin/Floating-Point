"""Local layout repairs after rendering the filled student report."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

p=Path(__file__).resolve().parents[1]/'Floating_Point_Report_Q0.docx'
d=Document(p)
for para in d.paragraphs:
    if para.text=='0 2 浮點型態資訊':para.paragraph_format.page_break_before=False
    if para.text.startswith('原敘述不夠精確。'):para.text=para.text.replace('[5]','[5b]')
    if para.text.startswith('本機 float／double 為 binary32') and 'binary32 的位元欄位' not in para.text:
        para.text += ' binary32 的位元欄位為 1 sign／8 exponent／23 fraction，binary64 為 1／11／52；normal 的隱含首位使有效位數為 24／53。x86 extended precision 的有效數首位明確儲存，為 1 sign／15 exponent／64 significand，共 80 位；此 ABI 的 16-byte 空間包含填補。'
    if para.text.startswith('圖 1-3 ') or para.text.startswith('圖 1-4 '):
        previous=para._p.getprevious()
        for inline in previous.xpath('.//wp:inline'):
            extent=inline.find(qn('wp:extent'));oldw=int(extent.get('cx'));oldh=int(extent.get('cy'));nw=int(Inches(5.6));nh=round(oldh*nw/oldw)
            extent.set('cx',str(nw));extent.set('cy',str(nh))
            for ext in inline.xpath('.//a:xfrm/a:ext'):ext.set('cx',str(nw));ext.set('cy',str(nh))
    if para.style.name=='Heading 2' and para.text=='1 7 小結':
        para.paragraph_format.space_before=Pt(6)
for section in d.sections:
    footer=section.footer.paragraphs[0]
    footer.text='Advanced C++ Homework 2　｜　'
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for run in footer.runs:run.font.size=Pt(9)
d.save(p)
print('Repaired page flow, figure sizes and footer.')
