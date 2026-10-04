"""One-time content migration from the completed 27-page report.

Already applied; the guard prevents applying it twice. Subsequent native DOCX
layout edits and visual QA produced the final 14-page report. This is an editing
record, not an experiment runner or a general report rebuild command.
"""
from pathlib import Path
import json,re,csv
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];file=ROOT/'Floating_Point_Report_Q0.docx'
d=Document(file);p=list(d.paragraphs);T=list(d.tables)
assert len(p)==200 and p[179].text.startswith('附錄 A'),'Requires the completed 27-page version'
# Original data remain in results, original screenshots remain in their folders.
remove=set([5,6,7,8,18,19,20,25,26,29,30,31,34,45,48,51,54,58,61,62,65,67,68,70,71,80,81,82,85,86,87,88,89,91,92,101,102,103,106,110,111,119,120,125,128,129,132,136,144,145,147,148,149,151,153,158,159,162,163,165,176,177,178])
# Remove whole oversized evidence blocks and duplicated setup screenshots.
remove.update([14,15,16,17])
# Keep Mac environment alongside the compact environment table; no separate page.
# Cut appendices from Word only; all raw files stay unchanged.
body=d._element.body;start=p[179]._p;idx=list(body).index(start)
for el in list(body)[idx:]:
 if el.tag!=qn('w:sectPr'):body.remove(el)
# Q5 Windows four-mode screenshot retained; Mac result captured in assembly and table.
replacements={
4:'我用執行輸出、計時與反組譯回答：這台機器如何處理浮點數，以及我能用什麼證據確認。結果顯示：float 與 double 不存在固定速度比例；GCC 的 fast-math 能改變結果；相同 long double 程式在 WSL 與 Mac 得到不同答案。以下依 Q0～Q7 說明實驗、證據與原因。',
10:'平台與重現條件',
12:'測試日期為 2026/10/04。WSL 的 12 核心／24 執行緒是虛擬拓樸；未固定實體 P/E 核心、即時頻率或溫度，因此 Q1 是這次執行的相對比較。Mac 僅作 Q5 對照，未量測跨平台速度。',
13:'重現：Windows 執行 experiments/q0_environment/collect_windows.ps1；WSL 執行 bash scripts/run_all.sh；Mac 執行 CXX=g++-16 bash scripts/run_q5.sh run。編譯命令、版本、來源雜湊、runtime 與輸出均由腳本保存。',
21:'型態參數也是實驗結果',
22:'q0-types.txt 實測 float／double／long double 的有效位數為 24／53／64。下表說明：sizeof=16 的 long double 並非 binary128，而是此 ABI 為 80-bit extended precision 保留的儲存空間。Mac 的 long double 則實測為 8 bytes／53 位，詳見 Q5。',
24:'WSL 同時量得 FLT_EVAL_METHOD=0、FE_TONEAREST、fast-math=false；epsilon 為 2^-23／2^-52／2^-63。__float128 僅偵測大小，未加入本次 benchmark。',
36:'目的與方法',
37:'我比較三型態在相同 scalar kernel 的加減、乘法與除法，分別量測一條相依鏈與四條獨立鏈。前者下一步依賴前一步結果，後者允許處理器重疊運算。',
38:'如何讓計時量到實際運算',
39:'每組暖機 2 次，再做 7 回合，每回合 5,000,000 迭代；一鏈／四鏈每迭代為 2／8 次算術運算。seed=20261004，初值於執行期產生，factor=1+1/1024；用 steady_clock 計時，輪換型態順序並取中位數。',
40:'kernel 分開編譯、設 noinline，關閉 LTO、fast-math、FMA contraction 與自動向量化；迴圈前後設 barrier，輸出 checksum。反組譯仍有算術迴圈與分支，證明計算未被消去。ns/op 包含迴圈及呼叫成本，不能直接當成硬體單條指令 latency。',
41:'來源：experiments/q1_benchmark/main.cpp、kernels.cpp。\n編譯：bash experiments/q1_benchmark/build.sh\n執行：./build/q1-scalar/q1_benchmark --iterations 5000000 --repeats 7 --warmup 2 --seed 20261004\n旗標：-std=c++20 -O3 -fno-fast-math -ffp-contract=off -fno-lto -fno-tree-vectorize',
42:'速度結果與證據',
43:'下表為 18 組結果的中位數，單位 ns/op＝總秒數×10^9÷算術運算次數。每組 7 回合共 126 筆，完整樣本與最小／最大值保留於 q1-timing-and-precision.txt、q1-summary.csv。',
49:'這台機器為什麼出現這些差異',
50:'單鏈乘法三型態皆約 0.86 ns/op，沒有隨型態大小形成固定倍數；單鏈除法則為 float 2.333、double 3.000、long double 3.327 ns/op。四鏈除法降為 0.618／0.828／0.923 ns/op，說明獨立運算能提高有效吞吐量。指令證據顯示 float／double 走 scalar SSE、long double 走 x87；精度、指令路徑、搬移與相依關係共同影響速度，不能只以 bytes 推算。[4]',
52:'精度實驗與結果',
53:'加總 10,000 組 [2^24,1,-2^24] 或 [2^53,1,-2^53]，分別正序與反序。所有輸入可精確表示，數學答案都是 10,000；用此整數作參考，不假定 long double 為真值。執行同一程式加 --precision-only。',
57:'圖 1-2 加總順序與型態改變結果。畫面命令指向本次結果目錄，與原始 log 一致。',
59:'從數值表示連結到硬體',
60:'binary32／binary64 的 sign、exponent、fraction 欄位為 1／8／23 與 1／11／52，normal 隱含首位使有效位數為 24／53。WSL long double 為 1 sign、15 exponent、64-bit significand 的 extended precision，共 80 位，ABI 儲存空間含填補。反組譯中 float 用 ADDSS／MULSS／DIVSS，double 用 ADDSD／MULSD／DIVSD，long double 用 x87 FADD／FMUL／FDIV；SSE 暫存器不代表本測試已向量化。',
66:'加減需對齊指數、運算有效數、正規化並捨入；乘除也只能保留有限精度。這解釋了 2^24 的 float、2^53 的 double 正序都丟失 +1，反序卻得 10,000。long double 的 64 位保住兩案例的小增量；float 在 2^53 兩順序都失去它。改順序有時改善結果，仍需按值域與誤差需求選擇演算法。指令可證明此次採用的路徑，無法揭露 CPU 內部全部微架構。[3][4]',
72:'我選擇 std::sqrt。C++20 浮點 overload 回傳 float、double 或 long double，整數呼叫等效使用 double，程式以 static_assert 核對。函式求非負平方根；負值的 domain error 依 math_errhandling 與實測判斷。[1][5]',
73:'比較 builtin 與函式庫路徑',
74:'// experiments/q2_cmath/sqrt_paths.cpp\nfloat sqrt_path(float x) { return std::sqrt(x); }\ndouble sqrt_path(double x) { return std::sqrt(x); }\nlong double sqrt_path(long double x) { return std::sqrt(x); }\n// 實際函式加 noinline；forced-libm 以 volatile function pointer 呼叫 symbol。',
75:'編譯：bash scripts/build_all.sh；執行：./build/all/q2_cmath 2 與 ./build/all/q2_libm 2。兩版皆為 -O2 -fno-fast-math -frounding-math -ffp-contract=off -fno-lto，libm 版另加 -DHW2_FORCE_LIBM=1。',
76:'本機究竟如何求平方根',
77:'反組譯顯示 default 版對三型態分別執行 SQRTSS、SQRTSD、FSQRT；負值跳往 sqrtf／sqrt／sqrtl。forced-libm 的 finite implementation 也找到硬體 sqrt 指令，因此本次呼叫函式庫仍可走硬體路徑。30 個輸入／型態組合的兩路徑輸出與例外旗標逐項一致；sqrt(4)=2 且精確，sqrt(2) 設 inexact，sqrt(-1) 回 NaN、errno=33、invalid=true，sqrt(-0) 保留負號。這些輸出驗證數值與錯誤處理，組合語言則證明採用的指令。',
83:'規格與實作能證明到哪一層',
84:'sqrt(2) 的三型態輸出如下，額外位數反映型態精度。最小 double subnormal 2^-1074 的平方根為 2^-537，實測未設 inexact；極小輸入不必然不精確。',
90:'Q3 IEEE 754 與 ffast math',
93:'關閉 fast-math 時，本例保留括號順序與特殊值語意；開啟後 GCC 允許重結合、假定只有有限值，並放寬 errno、signed zero 與捨入環境等要求。IEEE 754 的有限精度捨入使加法不具一般結合律，因此這些轉換可能改變答案。[2]',
94:'只改一個旗標的實驗',
95:'以 argv 讀入 a=1e16、b=-1e16、c=1，比較 (a+b)+c 與 a+(b+c)。運算置於獨立檔案、關閉 LTO，避免常數折疊。兩版皆 -std=c++20 -O3 -fno-lto -ffp-contract=off，只將 -fno-fast-math 換成 -ffast-math。另測 quiet NaN 的 x!=x 作特殊值補充案例。',
96:'來源：experiments/q3_fast_math/main.cpp、operations.cpp。\n編譯：bash scripts/build_all.sh\n執行：./build/all/q3_strict 1e16 -1e16 1\n      ./build/all/q3_fast 1e16 -1e16 1',
97:'結果與產生碼證據',
107:'結果如何被改變',
108:'strict 左結合先抵消為 0，再加 1 得 1；右結合的 -1e16+1 捨入回 -1e16，結果為 0。fast 的組合語言把左結合也改成先加 b+c，輸出便成 0；輸入、source 與其他旗標一致，足以把差異連結到 fast-math。NaN 的 x!=x 在 strict 為 true，fast 則直接以 XOR 回 false，說明有限值假設會破壞這類 NaN 檢查。兩版皆禁用 contraction，這次差異不需要用 FMA 解釋。[2]',
112:'反例與實驗',
113:'「固定輸出 9 位小數」不等於「十進位中點一律進位」。我用 1/1024=0.0009765625，此值可精確以二進位表示，正好在第 9 位的中點；因而可排除字面值轉換誤差。比較 printf、iostream、正負值、nextafter 鄰居與四種捨入模式。',
116:'nearest 時兩 API 都輸出 0.000976562，與中點一律進位的 0.000976563 不同；較小／較大鄰居分別輸出 0.000976562／0.000976563。四模式結果如下，兩 API 在 16 個測試組合一致。',
121:'這台機器的 nearest 輸出符合 ties-to-even，upward／downward 又顯示格式化受捨入環境影響。因此應說「把實際儲存值格式化為指定小數位」，不能無條件宣稱是原始十進位數的四捨五入；也不將本次函式庫行為擴大為所有平台的保證。[5b]',
123:'跨平台條件與程式',
124:'我在 x86_64 WSL 與 arm64 Mac 以相同來源、GNU GCC release 16.2.0 及逐字相同旗標計算 (2^53+1)-2^53。兩端都先設定 FE_TONEAREST；pair checker 對 compiler family、release、flags、Q5／common 來源雜湊皆 MATCH。GCC 的 vendor build 與 target 不同，完整 configuration 均保留。',
126:'輸出與指令是否支持 ABI 差異',
127:'// experiments/q5_portability/operations.cpp；獨立函式加 noinline\nlong double increment_long_double(long double a, long double b) {\n    long double sum = a + b;\n    asm volatile("" : "+m"(sum) : : "memory");\n    return sum - a;\n}\n// 輸入 a=2^53、b=1；memory barrier 強制中間值存回平台型態。',
133:'編譯／執行：CXX=g++-16 bash scripts/run_q5.sh run。旗標為 -std=c++20 -Wall -Wextra -fno-lto -ffp-contract=off -O2 -fno-fast-math -frounding-math。用 scripts/verify_q5_pair.py 比較兩端結果目錄；圖 5-3 為核對輸出。',
141:'WSL 函式用 FLDT、FADD、FSTPT、FLDT、FSUBP；Mac 用 FADD D、STR／LDR、FSUB D。兩端都保留加法、存回與減法，因此差異不是計算被消去。輸入有限且不溢位，沒有越界或未初始化值；inline asm 是此次 GCC 的支援擴充。',
142:'同一計算為何一邊得 1 一邊得 0',
143:'WSL long double 實測 16 bytes／64 位，能保留 2^53+1，減回得 1。Mac 實測 8 bytes／53 位，2^53 附近的間距為 2，+1 是中點，nearest 存回 2^53，減回便得 0；與 Apple arm64 的 long double ABI 一致。[6] 相同 source、release 和 flags 仍可能因型態表示而不同。本例另以同 binary 的 double 1+2^-53 對照四模式：兩端 nearest／downward／towardzero 得 1，upward 得下一個可表示值。捨入狀態与產生碼一起支持原因，而非僅憑理論推測。',
150:'這些實驗讓我把安全性視為誤差與重現要求：Q1 顯示需按值域選精度，速度要實測；Q2 顯示規格、函式庫与硬體是不同層次；Q3 顯示 fast-math 會改答案；Q4 顯示短輸出可能遮住儲存值與捨入規則；Q5 顯示 long double 名稱相同仍可能有不同精度。科學計算或碰撞判定應先訂誤差容忍度，再選演算法；比較值時依用途使用絕對／相對容差，診斷時保留 max_digits10 或 hexfloat。',
152:'我另測 gradual underflow，回答「結果是 subnormal 是否就一定觸發 underflow？」測三算式 min_normal×0.5、denorm_min×0.5、denorm_min×1，記錄分類與 FE_UNDERFLOW／FE_INEXACT。',
154:'來源：experiments/q6_extra/main.cpp、operations.cpp。\n編譯：bash scripts/build_all.sh；執行：./build/all/q6_extra。\nnearest 下讀到 MXCSR=0x1f80、FTZ／DAZ=false；未更動其設定。',
160:'三型態結果一致：精確 subnormal 不設 UF／inexact；最小 subnormal 再乘 0.5 捨入為零，兩旗標皆 true；乘 1 則保留。這證明「很小」「subnormal」「underflow exception」是不同判斷。接近零時相對精度下降，flush 設定也需注意；MXCSR 僅支持 SSE 狀態解釋，不能代替 x87 控制字。本次未量測 subnormal 速度。',
161:'Q7 AI 使用揭露',
164:'我提供設備資訊，在 Windows／WSL 執行 Q0～Q6、在 Mac 執行 Q5，並顯示結果後截圖。Codex 協助程式、保存腳本、工具鏈設定、來源查閱、流程預檢，以及數據整理與分析草稿。正式數據取自兩端 student 結果目錄；assistant 試跑未作為正式 benchmark。原因分析由 AI 協助撰寫，我仍需逐段確認理解。',
167:'參考文件用來解釋規格；本機執行路徑以輸出與反組譯確認。glibc generic master 的軟體分支會縮放、以 reciprocal sqrt 初值和多項式修正並補償捨入，但未與 Ubuntu 套件 source 精確匹配，不能當成本次 CPU 內部演算法。[5a]',
}
for i,s in replacements.items():p[i].text=s
for i in sorted(remove,reverse=True):
 el=p[i]._p
 if el.getparent() is not None:el.getparent().remove(el)
# Remove redundant Q5 control matrix and Q7 duplicated division-of-work table.
for i in [7,10]:
 if T[i]._tbl.getparent() is not None:T[i]._tbl.getparent().remove(T[i]._tbl)
# Single compact matrix retains every benchmark median.
def replace_table(old,headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));old._tbl.addprevious(t._tbl);old._tbl.getparent().remove(old._tbl)
 for j,row in enumerate([headers]+rows):
  cells=t.rows[0].cells if j==0 else t.add_row().cells
  for k,s in enumerate(row):cells[k].text=str(s)
 t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 return t
rs=list(csv.DictReader((ROOT/'results/linux-20261004-201349/q1-summary.csv').open()))
print('CSV fields',rs[0].keys())
# Use preexisting table cells as a verified exact source for display values.
vals={}
for row in T[2].rows[1:]:
 c=[x.text for x in row.cells];vals[c[0],c[1],c[2]]=c[3]
rows=[]
for op in ['add_sub','multiply','divide']:
 for chain in ['1','4']:rows.append([op,'相依鏈' if chain=='1' else '四獨立鏈']+[vals[op,chain,typ] for typ in ['float','double','long double']])
replace_table(T[2],['運算','模式','float','double','long double'],rows,[1.1,1.3,1.3,1.3,1.5])
# Retain all accuracy results in six paired rows.
r=[]
for k in range(1,13,2):
 a=[c.text for c in T[3].rows[k].cells];b=[c.text for c in T[3].rows[k+1].cells]
 r.append([a[0],a[1],a[3],b[3],a[4]+'／'+b[4]])
replace_table(T[3],['型態','大數','正序結果','反序結果','絕對誤差 正／反'],r,[1.3,0.7,1.1,1.1,2.3])
# Condense Q6 repeated type records: all three share the same exception pattern.
replace_table(T[9],['算式','三型態結果分類','UF／inexact'],[['min_normal × 0.5','subnormal（精確）','false／false'],['denorm_min × 0.5','zero','true／true'],['denorm_min × 1','subnormal（保留）','false／false']],[2.2,2.3,2.0])
# Shorten captions; retain figure number so raw evidence can still be located.
for x in d.paragraphs:
 if x.text.startswith('圖 '):
  s=x.text.split('（原檔：')[0].strip().replace('本人的','').replace('本人','').replace('本次','')
  x.text=s
  for run in x.runs:run.font.size=Pt(9)
 if x.style.name.startswith('Heading') and x.style.name!='Heading 1':x.paragraph_format.page_break_before=False
 if x.text and x.style.name=='Normal':
  x.paragraph_format.space_after=Pt(5);x.paragraph_format.space_before=Pt(0)
 for phrase in ['本人實測','本人採集','本人 Windows／WSL','本人 Mac']:
  if phrase in x.text:x.text=x.text.replace(phrase,phrase.replace('本人',''))
# Strip actor labels from Q0 and Q5 table columns.
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   if '本人' in c.text:c.text=c.text.replace('本人','')
# Condense environment rows without losing reproducibility information.
env=d.tables[0]
env.cell(7,2).text='Q5 同旗標（見 Q5）';env.cell(12,1).text='平衡／AC／100%\n2026/10/04';env.cell(12,2).text='2026/10/04\n電源未採集'
# Moderate picture widths; retain original media bytes and any native crops.
for x in d.paragraphs:
 if x.text.startswith('圖 '):
  prev=x._p.getprevious()
  if prev is None:continue
  for inline in prev.xpath('.//wp:inline'):
   ex=inline.find(qn('wp:extent'));ow=int(ex.get('cx'));oh=int(ex.get('cy'));width=5.2
   if x.text.startswith(('圖 0-4','圖 1-2','圖 3-3','圖 5-3')):width=6.5
   if x.text.startswith('圖 2-1'):width=5.4
   nw=int(Inches(width));nh=round(oh*nw/ow);ex.set('cx',str(nw));ex.set('cy',str(nh))
   for ee in inline.xpath('.//a:xfrm/a:ext'):ee.set('cx',str(nw));ee.set('cy',str(nh))
# Compact references retain original support links.
for x in d.paragraphs:
 if re.match(r'^\[\d',x.text):
  lines=x.text.splitlines();x.text=lines[0]+'\n'+lines[1]
  for run in x.runs:run.font.size=Pt(9)
# Unified table spacing and repeated headers, readable 10pt.
for t in d.tables:
 for j,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if j==0:pr.append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   for x in c.paragraphs:
    x.paragraph_format.space_before=Pt(2);x.paragraph_format.space_after=Pt(2)
    for run in x.runs:run.font.size=Pt(10)
# One short repository pointer replaces raw-data appendices.
x=d.add_paragraph('原始資料與重現檔案',style='Heading 2')
d.add_paragraph('WSL：results/linux-20261004-201349/；Mac：results/q5-darwin-20261004-211048-student-wdBjAR/。完整計時、案例、來源雜湊與組合語言保留於 results；截圖保留於 screenshot／mac截圖。各題 experiments README 提供獨立重現步驟。')
d.save(file)
m=ROOT/'docs/student_report_evidence.json';data=json.loads(m.read_text(encoding='utf-8'))
labels={x.text.split(' ')[1] for x in d.paragraphs if x.text.startswith('圖 ')}
for f in data['figures']+data['mac_figures']:f['included_in_condensed_report']=f['caption'].split(' ')[1] in labels
for k in ['report_pages','visual_qa']:data.pop(k,None)
data['embedded_screenshots']=len(d.inline_shapes);data['report_focus']='machine observations, execution evidence, explanations; AI disclosure in Q7 only';m.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('Condensed report:',len(d.inline_shapes),'figures,',len(d.paragraphs),'paragraphs')
