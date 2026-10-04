"""Fill the existing report only from the student's saved run and screenshots."""
from pathlib import Path
import csv, json, hashlib, re, statistics
from collections import defaultdict
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
RUN=ROOT/'results/linux-20261004-201349'
DOC=ROOT/'Floating_Point_Report_Q0.docx'
assert 'executed_by=student' in (RUN/'start.txt').read_text()
assert 'failed_experiments=0' in (RUN/'completion.txt').read_text()
d=Document(DOC); p=list(d.paragraphs); oldtables=list(d.tables)
assert len(p)==121 and p[5].text=='填寫方式', 'This one-time fill script requires the preparation template; refuse to overwrite a filled report.'

def put(i,s): p[i].text=s
def extra(i,s,style=None):
    x=p[i].insert_paragraph_before(s,style=style)
    return x
def fmt_table(t,widths=None,size=10):
    t.autofit=False
    for j,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        for h in list(trpr.findall(qn('w:trHeight'))):trpr.remove(h)
        el=OxmlElement('w:cantSplit');trpr.append(el)
        if j==0:
            el=OxmlElement('w:tblHeader');trpr.append(el)
        for k,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[k])
            pr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2' if j==0 else ('F7F8FA' if j%2==0 else 'FFFFFF'));pr.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3)
                para.paragraph_format.keep_with_next=False
                for r in para.runs:r.font.size=Pt(size);r.bold=(j==0)
    if widths:
        for col,wid in zip(t.columns,widths):col.width=Inches(wid)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+name);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
def table(headers,rows,anchor=None,replace=None,widths=None,size=10):
    t=d.add_table(rows=1,cols=len(headers))
    for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
    for vals in rows:
        for c,s in zip(t.add_row().cells,vals):c.text=str(s)
    fmt_table(t,widths,size)
    if replace is not None:
        replace._element.addprevious(t._element);replace._element.getparent().remove(replace._element)
    elif anchor is not None:anchor._p.addprevious(t._element)
    return t

figures=[]
def fig(i,name,caption,crop=None,width=6.5):
    f=ROOT/'screenshot'/name
    x=p[i].insert_paragraph_before()
    shape=x.add_run().add_picture(str(f),width=Inches(width))
    iw,ih=Image.open(f).size
    if crop:
        left,top,right,bottom=crop
        fill=shape._inline.xpath('.//pic:blipFill')[0]
        rect=OxmlElement('a:srcRect')
        for k,v in [('l',left/iw),('t',top/ih),('r',(iw-right)/iw),('b',(ih-bottom)/ih)]:rect.set(k,str(round(v*100000)))
        fill.insert(1,rect)
        shape.height=Inches(width*(bottom-top)/(right-left))
    x.paragraph_format.keep_with_next=True;x.paragraph_format.space_after=Pt(4)
    c=p[i].insert_paragraph_before(caption+'（原檔：'+name+'）')
    c.paragraph_format.space_after=Pt(10)
    for r in c.runs:r.font.size=Pt(9)
    figures.append({'file':name,'caption':caption,'crop':crop,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
def block_cases(name,prefix='type='):
    s=(RUN/name).read_text(encoding='utf-8')
    blocks=[]
    current=None
    for line in s.splitlines():
        if line.startswith(prefix):
            current=dict(x.split('=',1) for x in line.split(','));blocks.append(current)
        elif current is not None and '=' in line and not line.startswith(('$','[exit')):
            for f in line.split(','):
                if '=' in f:
                    k,v=f.split('=',1);current[k]=v
    return blocks

put(5,'資料來源與完成範圍')
put(6,'本報告整理我於 2026/10/04 親自執行的 Windows／WSL 實驗與截圖。Windows 採集時間為 20:07:12；WSL 批次於 20:13:49 開始，結果目錄為 results/linux-20261004-201349，executed_by=student，failed_experiments=0。所有數據由此批次讀取；助手先前的試跑不混入表格。Mac Q5 的本人重跑與截圖尚待補入。')
env=oldtables[0]
env.cell(0,1).text='主要平台（本人實測）'
updates={3:('x86_64-pc-linux-gnu','aarch64-apple-darwin27（先前 Mac 助手紀錄）'),4:('Windows 11 家用版 25H2\nBuild 26200.9457','macOS 27.0.1，build 26A434\n先前助手採集；本人 Q5 待補'),6:('GNU GCC 16.2.0\nHW2 local GCC 16.2.0','GNU GCC 16.2.0\nHomebrew；本人重跑待補'),7:('-std=c++20 -Wall -Wextra -fno-lto\n-ffp-contract=off -O2\n-fno-fast-math -frounding-math','Q5 將使用相同旗標'),11:('glibc 2.43-2ubuntu2.4\n本地 GCC 16.2.0 的 libstdc++／libgcc_s','本人實際連結紀錄待補'),12:('平衡；AC 電源；電池 100%\n本人採集：2026/10/04 20:07','裝置資訊由本人提供\n本人 Q5 執行時間／電源待補')}
for i,(a,b) in updates.items():env.cell(i,1).text=a;env.cell(i,2).text=b
fmt_table(env,size=10)
put(12,'WSL 的 12 核心／24 執行緒是虛擬拓樸，不能視為實體 P/E 核心配置。實驗使用 WSL 使用者目錄中的 GCC 16.2.0；Windows 路徑下的 llvm-mingw 與系統 GCC 15.2.0 未用於本次數據。CPU 即時頻率、溫度與實體核心排程未控制，這是效能比較的限制。')
put(13,'重現指令：Windows PowerShell 執行 ./experiments/q0_environment/collect_windows.ps1；WSL 專案根目錄執行 bash scripts/run_all.sh。腳本保存編譯命令、輸出、退出碼、旗標、來源 SHA-256、組合語言與函式庫依賴。')
fig(14,'Screenshot 2026-10-04 200830.png','圖 0-1 本人的 Windows 硬體、WSL 與電源採集。Word 只裁去右側空白，原始截圖完整保存。',(0,0,1130,1439),width=5.4)
fig(14,'Screenshot 2026-10-04 201929.png','圖 0-2 Ubuntu 26.04.1、WSL kernel、x86_64 與 i7-14650HX。')
fig(14,'Screenshot 2026-10-04 201656.png','圖 0-3 批次完成且 failure=0，實際 compiler 為 GCC 16.2.0；顯示左側版本資訊局部。',(0,0,1100,150))
put(14,'')
put(16,'下表來自本次 GCC 16.2.0 的 q0-types.txt。sizeof 是 ABI 儲存大小，digits 才是二進位有效位數；long double 的 16 bytes 並不代表 binary128。')
fmt_table(oldtables[1],size=10)
put(19,'單獨重現 Q0：先 source scripts/compiler_env.sh，再以 "$CXX" -std=c++20 -O2 -fno-fast-math -ffp-contract=off experiments/q0_environment/main.cpp -o build/q0_environment 編譯，執行 ./build/q0_environment。本次完整批次另加 -Wall -Wextra -fno-lto -frounding-math，精確命令見 build.txt。')
put(20,'本人原始資料：results/q0-windows/20261004-200712/、results/linux-20261004-201349/q0/environment.txt 與 q0-types.txt。')
fig(21,'Screenshot 2026-10-04 201840.png','圖 0-4 本次型態輸出的左側局部，包含 sizeof、alignment、digits 等；完整欄位保留於原圖與 q0-types.txt。',(0,0,1100,265))
put(21,'');put(22,'本次型態表與正式數據均使用 GCC 16.2.0。__float128 只偵測大小，未做其計算或效能比較。')

rows=list(csv.DictReader((RUN/'q1-summary.csv').open(encoding='utf-8')))
assert len(rows)==18
def six(x):return f'{float(x):.6f}'
table(['運算','鏈數','型態','中位數 ns/op','最小 ns/op','最大 ns/op'],[[r['operation'], '1' if r['mode']=='latency1' else '4',r['type'].replace('_',' '),six(r['median_ns_per_operation']),six(r['min_ns_per_operation']),six(r['max_ns_per_operation'])] for r in rows],replace=oldtables[2],widths=[1,0.45,1,1.35,1.35,1.35],size=9.5)
put(25,'以相同 scalar kernel 比較 float、double 與 long double 的計算速度與精度。測量加減、乘法、除法三組操作，並區分相依鏈與四條獨立鏈；本次不包含 __float128。')
put(27,p[27].text.replace('準備測量','實際測量'))
put(31,'以下為本人實際執行的 18 組摘要，每組 7 次，共 126 筆計時。ns/op = seconds × 10^9 ÷ 算術運算次數；一條鏈每迭代 2 次，四條鏈每迭代 8 次。數字四捨五入顯示至 6 位小數；原始精度與各回合資料保留於 log，附錄也列出各回合 ns/op。')
put(33,'一條鏈較偏向相依運算成本；四條鏈讓處理器重疊獨立運算。這些時間含迴圈控制、呼叫與少量收尾成本，不應直接標成單一指令 latency。')
fig(34,'Screenshot 2026-10-04 202051.png','圖 1-1 本人產生的全部 Q1 中位數摘要與迭代設定。',width=5.1);put(34,'')
put(36,'本機 float／double 的加減與乘法接近，long double 的加減较慢，除法則 float 最快、double 次之、long double 最慢。單鏈乘法三型態中位數約 0.86 ns/op，顯示較高精度不必然在每一種算術測試都較慢。四鏈除法約為 0.618／0.828／0.923 ns/op，而單鏈為 2.333／3.000／3.327 ns/op；可見獨立鏈能提高有效吞吐量。因沒有固定 P/E 核心、頻率、溫度或背景負載，這是本次執行的相對觀察，不能宣稱固定速度比例。')
acc=[]
for line in (RUN/'q1-timing-and-precision.txt').read_text().splitlines():
    if line.startswith('accuracy,'):
        f=line.split(',');v=dict(x.split('=',1) for x in f[4:]);acc.append([f[1].replace('_',' '),f[2].split('=',1)[1],f[3],v['result'],v['absolute_error'],v['relative_error']])
assert len(acc)==12
table(['型態','大數','順序','結果','絕對誤差','相對誤差'],acc,replace=oldtables[3],size=10)
fig(42,'Screenshot 2025-07-15 221151.png','圖 1-2 本次 12 個精度案例。檔名含舊日期，但畫面指令指向 linux-20261004-201349，逐行與該 log 一致；不以檔名推定實驗日期。');put(42,'')
put(44,'本機 float／double 為 binary32／binary64 特性，有效位數分別為 24／53；long double 為 x86 extended precision 特性，有效位數 64、儲存大小 16 bytes。kernels.s 中 float 使用 ADDSS／SUBSS／MULSS／DIVSS，double 使用 ADDSD／SUBSD／MULSD／DIVSD，long double 使用 x87 FADD／FSUB／FMUL／FDIV。自動向量化已關閉，SSE 暫存器的使用不代表本測試一次處理多個數。')
fig(45,'Screenshot 2026-10-04 202216.png','圖 1-3 float kernel 的 scalar SSE 算術指令。')
fig(45,'Screenshot 2026-10-04 202254.png','圖 1-4 double kernel 的 scalar SSE 算術指令。');put(45,'')
put(46,'硬體算術必須先處理表示格式與符號。加減需要對齊指數、處理有效數相加減、正規化並捨入；乘法結合指數與有效數乘積；除法產生有限精度商並捨入。一般理解模型中的 guard／round／sticky 資訊用來判斷捨入，不能據此聲稱本 CPU 的內部微架構細節。每條相依鏈的後一次計算需要前一次結果，四條獨立鏈則可由排程器重疊執行；x87 的 stack register／資料搬移成本也與 SSE 路徑不同。實際指令與迴圈仍存在的證據在 kernels.s 與 q1-kernels.disasm.txt，未憑型態名稱推論速度。')
put(48,'精度結果顯示，所有輸入即使可精確表示，中間加法仍可能捨入。2^24 案例的 float 正序結果為 0，反序為 10000；2^53 案例的 double 也出現此順序差異，long double 兩順序均為 10000。float 在 2^53 案例兩順序均失去小增量。因此改變順序有時能改善結果，但不是通用保證；需要依輸入值域與誤差需求選型態、演算法與加總方式。')

q2a=block_cases('q2-default.txt');q2b=block_cases('q2-libm.txt')
assert len(q2a)==len(q2b)==30
keys=['sqrt','errno','FE_INVALID','FE_INEXACT','negative_sign','isnan']
assert all(all(a[k]==b[k] for k in keys) for a,b in zip(q2a,q2b))
put(57,'本機 default assembly 對 float／double／long double 分別有 SQRTSS、SQRTSD、FSQRT；遇負值另跳往 sqrtf／sqrt／sqrtl，保留 domain error 路徑。forced-libm 透過 volatile function pointer 進入實際 symbol，函式庫 finite implementation 的反組譯也有相同硬體 sqrt 指令，故呼叫 libm 並不等於採軟體近似。glibc 的 generic 參考來源有 builtin 分支，另有軟體分支：正規化 exponent／mantissa，以 reciprocal sqrt 表格初值及多項式修正，再補償誤差處理捨入，極小值先縮放。這是參考演算法，master 尚未與本機套件原始碼精確匹配；本機硬體路徑以實際反組譯為證。[5]')
fig(58,'Screenshot 2026-10-04 202449.png','圖 2-1 default 的硬體 sqrt／負值 fallback，以及 forced-libm 的符號載入。')
fig(58,'Screenshot 2026-10-04 202517.png','圖 2-2 實際連結本地 GCC runtime 與系統 libm；glibc 版本為 2.43。');put(58,'')
table(['型態','default 指令／fallback','sqrt(2) 實測'],[[a['type'].replace('_',' '),{'float':'SQRTSS／sqrtf','double':'SQRTSD／sqrt','long_double':'FSQRT／sqrtl'}[a['type']],a['sqrt']] for a in q2a if a['case']=='runtime'],replace=oldtables[4],widths=[1,2.5,3],size=10)
extra(62,'兩版本的 30 個案例在回傳值、errno、invalid／inexact 與符號上逐筆一致。sqrt(4) 皆為 2 且 exact；sqrt(2) 設定 inexact；負值回傳 NaN、errno=33（EDOM）、invalid=true；負零保持負號。最小 subnormal 的平方根不必然 inexact，例如 double 的 2^-1074 開根號為可精確表示的 2^-537。完整三型態案例表見附錄。')
fig(62,'Screenshot 2026-10-04 202419.png','圖 2-3 double 特殊值、平方根結果與 errno／exception 的實測。')
put(62,'本次 math_errhandling=3，同時支援 errno 與浮點例外。std::sqrt 的規格、compiler 內聯策略、函式庫 wrapper 與硬體指令是不同層次；看到 SQRTSD 只能證明該指令路徑，不能猜測 CPU 內部平方根使用何種迭代演算法。')

table(['觀察','strict','fast'],[['(a+b)+c','1','0'],['a+(b+c)','0','0'],['NaN 的 x!=x','true','false'],['finite control','6','6']],replace=oldtables[5],size=10)
fig(72,'Screenshot 2026-10-04 202616.png','圖 3-1 相同 a、b、c 輸入下 strict／fast 的完整輸出。');put(72,'')
fig(73,'Screenshot 2026-10-04 202736.png','圖 3-2 本機 -Q 記錄：重結合與有限值假設等子旗標的差異。')
fig(73,'Screenshot 2026-10-04 202817.png','圖 3-3 ADD 加法順序改變，fast 版也移除 NaN self comparison。');put(73,'')
put(75,'strict 的左結合先讓 1e16 與 -1e16 抵消為 0，再加 1 得 1；右結合中 -1e16+1 捨入回 -1e16，最後得 0。fast 的 left_grouped assembly 先加 b+c，因而同樣得 0，這個有限值案例已證明僅改旗標便改變結果。NaN 的 x!=x 在 strict 為 true，fast 的 self_inequality 則以 xorl 直接回 false，展示 finite-math-only 的假設；此補充案例不作為一般 NaN 檢查的安全寫法。-ffast-math 並非提升硬體精度，而是放寬必須保留的數值語意；本次還觀察到 associative／unsafe-math／finite-math-only 啟用，math-errno／signed-zeros／trapping-math 停用。兩版同關閉 contraction，差異無須以 FMA 解釋。[2]')

put(80,p[80].text.replace('結果仍待本人驗證。','本次以 nearest、upward、downward、towardzero 逐一實測。'))
table(['模式','正中點：兩 API','負中點：兩 API'],[['nearest','0.000976562','-0.000976562'],['upward','0.000976563','-0.000976562'],['downward','0.000976562','-0.000976563'],['towardzero','0.000976562','-0.000976562']],replace=oldtables[6],size=10)
extra(83,'兩 API 在本次所有案例一致。nearest 下，中點的較小鄰居 0x1.fffffffffffffp-11 輸出 0.000976562，較大鄰居 0x1.0000000000001p-10 輸出 0.000976563。四模式與正負值／鄰居完整 16 個組合見附錄。')
fig(83,'Screenshot 2026-10-04 203035.png','圖 4-1 精確中點、鄰近值與四模式的關鍵輸出。',width=5.2);put(83,'')
put(85,'原敘述不夠精確。如果「四捨五入」表示中點一律進位，0.0009765625 應為 0.000976563；但本機 nearest 的 printf 與 iostream 實際為 0.000976562，符合保留末位偶數的中點處理。此數本身可精確二進位表示，所以反例不是字面值轉換誤差造成。格式化的是實際儲存值，且本機輸出受 rounding mode 影響；這一輪結果不能單独當成所有平台／C++ implementation 的格式化保證。[5]')

put(88,'本節的 Windows／WSL 部分已由本人完成。A 使用同一 binary、同一數值輸入，只更動 rounding environment；B 準備用 GNU GCC 16.2.0、同 source／flags，在兩平台 nearest 下比較 long double ABI。Mac 先前助手測試可用於流程預檢，但本報告的本人跨平台結論暫不完成，等本人在 Mac 重跑並截圖後補入。')
src=(RUN/'source-sha256.txt').read_text()
q5hash='\n'.join(line for line in src.splitlines() if 'experiments/q5_portability/' in line or 'experiments/common/' in line)
table(['控制條件','本人 Windows／WSL','本人 Mac'],[['來源','Q5／common SHA-256 已保存\n見 source-sha256.txt','待本人跑後用 pair checker 核對'],['compiler','GNU GCC 16.2.0\nHW2 local build','預定 Homebrew GNU GCC 16.2.0'],['旗標',(RUN/'q5-flags.txt').read_text().strip(),'必須使用同旗標；待紀錄'],['target','x86_64-pc-linux-gnu','待本人執行採集'],['ABI／nearest','long double 16 bytes／64 digits\n(2^53+1)-2^53 = 1','待本人結果與截圖'],['OS／runtime','Ubuntu 26.04.1、glibc 2.43\n本地 GCC runtime','待本人實際依賴紀錄']],replace=oldtables[7],size=9.5)
put(91,p[91].text.replace('結果與條件仍待本人採集。','Windows 結果已採集；Mac 本人對照待補。'))
fig(92,'Screenshot 2026-10-04 203143.png','圖 5-1 本人 Q5 的 GCC release、逐字旗標與同 binary SHA-256。')
fig(92,'Screenshot 2026-10-04 203214.png','圖 5-2 同一 binary 的四種捨入模式與 nearest ABI 案例。');put(92,'')
put(93,'Mac 待補：在專案根目錄執行 CXX=g++-16 bash scripts/run_q5.sh run，保存本人 q5-nearest.txt、版本／flags／來源雜湊與截圖，再與本次 Windows 目錄執行 verify_q5_pair.py。不得以已保存的 Mac assistant 紀錄冒充本人執行。')
put(95,'候選 A：double 的 1+2^-53 是 1 與下一個 representable 值之間的中點。本人結果在 nearest／downward／towardzero 為 0x1p+0，upward 為 0x1.0000000000001p+0；同 binary SHA-256 可排除編譯變動。候選 B：WSL long double 有 64 位有效精度，能保留 2^53+1，實測減回 2^53 為 1。Apple 官方 arm64 ABI 規定 long double 等同 double，因此預期有效精度較小可能產生差異，但本人 Mac 結果與條件核對尚待完成。[6]')
put(97,'候選 A 的受控環境差異已由本人驗證；候選 B 暫不宣稱完成跨平台加分證據。相同 release 不表示 compiler binary／configuration／vendor patch 相同，需一併保存；GCC 對 frounding-math 的文件限制也使旗標本身不足以保證行為，應以此案例的產生碼與實測共同核對。[2]')

put(100,'本機結果支持按需求選擇浮點型態：float／double 不會僅因大小不同就有固定速度比例；較高精度能保留某些小增量，但仍是有限精度。std::sqrt 可能經 compiler builtin 或 libm，兩者最終都可使用硬體。fast-math 能改重結合與特殊值語意；十進位輸出也不能概括為中點一律進位。這些結論分別由 Q1～Q4 的數據與截圖支持，Q5 的本人 Mac 對照仍待完成。')
put(102,'科學計算或碰撞判定应先訂誤差容忍度與輸入值域，再選型態與演算法；對大量加總可評估較高精度或補償加總，但不能未測即保證改善。需要 reproducibility 時固定 source、compiler／flags、浮點環境與函式庫，避免任意重結合。浮點比較應依用途設定絕對／相對容差；精確可表示的控制值仍可合理使用 ==，不是一律禁止。輸出時保留 max_digits10 或 hexfloat 以便診斷，不讓短十進位顯示遮住差異。')
q6=block_cases('q6-underflow.txt')
table(['型態／案例','result / hexfloat','分類','UF／inexact'],[[a['type'].replace('_',' ')+'\n'+a['case'],a['result']+'\n'+a['hex'],a['result_class'],a['FE_UNDERFLOW']+'／'+a['FE_INEXACT']] for a in q6],replace=oldtables[8],widths=[1.7,2.9,0.8,1.1],size=9)
fig(108,'Screenshot 2026-10-04 203307.png','圖 6-1 float 的三個案例；nearest、FTZ／DAZ=false，subnormal 與 exception flags 已區分。',width=5.5);put(108,'')
put(110,'三型態皆觀察到：min_normal×0.5 是精確 subnormal，UF／inexact 均 false；denorm_min×0.5 在 nearest 捨入為零，兩旗標均 true；denorm_min×1 則保留 subnormal 且旗標均 false。因此「結果很小」、「是 subnormal」、「發生 underflow exception」不是同一件事。接近零時的相對有效精度下降與 flush 設定會影響數值判斷；本次 MXCSR=0x1f80，FTZ／DAZ=false，只能用於 SSE 環境解釋，不能直接代替 x87 控制狀態。本次沒有測 subnormal 速度，不宣稱其效能比例。')
put(113,'以下依目前實際分工揭露；本人 Mac Q5 完成後還需更新其執行紀錄。')
put(114,'我提供 Windows／Mac 設備與 compiler 資訊，親自在 Windows PowerShell／WSL 執行環境採集與完整 Q0～Q6 批次，並以終端顯示本人結果後截圖。Codex 協助建立程式、建置與保存腳本、安裝 WSL GCC 16.2.0、查閱來源、先做助手驗證，以及將我的真實數據與截圖整理成表格和分析。助手試跑目錄標為 assistant；本報告主要表格取自我執行的 linux-20261004-201349，不混用助手計時。原因分析文字由 AI 協助草擬，我仍需逐段驗證理解；本人 Mac Q5 尚待完成。')
table(['分工','實際完成内容'],[['本人','提供設備資訊；Windows／WSL 採集、完整執行與截圖；Mac 本人 Q5 待補'],['Codex','程式／腳本、工具鏈 setup、來源查閱、助手試跑、數據整理與分析草稿'],['資料界線','本人 student 目錄作主要證據；assistant 資料保留為準備驗證，未混入本人計時']],replace=oldtables[9],widths=[1.1,5.4],size=10)
put(117,'下列是規格、工具與參考實作來源。理論文件與本機實測分開：glibc master 是參考來源，未聲稱與 Ubuntu 套件 source 精確一致；本機路徑以反組譯佐證。查閱日期：2026/10/04。')
refs=[('[1] C++ working draft c.math','https://eel.is/c++draft/c.math','現行草案；C++20 浮點 overload 另由程式 static_assert 檢查。'),('[2] GCC 16.2.0 Optimize Options','https://gcc.gnu.org/onlinedocs/gcc-16.2.0/gcc/Optimize-Options.html','fast-math、contraction、rounding-math。'),('[3] glibc Floating Point Parameters','https://sourceware.org/glibc/manual/latest/html_node/Floating-Point-Parameters.html','型態與精度參數。'),('[4] Intel SDM 官方入口','https://www.intel.com/content/www/us/en/developer/articles/technical/intel-sdm.html','SSE／x87 指令來源入口；不由入口推論機型 latency 或內部演算法。'),('[5] glibc Exponents and Logarithms','https://sourceware.org/glibc/manual/latest/html_node/Exponents-and-Logarithms.html','sqrt 與 domain error。'),('[5a] glibc generic double e_sqrt.c','https://raw.githubusercontent.com/bminor/glibc/master/sysdeps/ieee754/dbl-64/e_sqrt.c','master 參考演算法，未與本機 source 精確匹配。'),('[5b] glibc Rounding Modes','https://sourceware.org/glibc/manual/latest/html_node/Rounding.html','捨入環境；格式化結果仍以本機輸出為證。'),('[6] Apple Writing ARM64 code','https://developer.apple.com/documentation/xcode/writing-arm64-code-for-apple-platforms','Apple ABI；本人 Mac Q5 待補。')]
put(118,'')
for title,url,note in refs:extra(118,title+'\n'+url+'\n'+note)
put(119,'目前完成與待補範圍')
put(120,'Windows／WSL 的本人數據、原始 log 與截圖已整理。待補本人 Mac Q5 的版本／flags／source／ABI 證據、實測輸出與截圖，再完成最後跨平台結論。提交時另輸出 PDF，命名為 113062331_林欣磊.pdf。')

# Detailed tables retain every student's measured timing sample and numerical case.
d.add_heading('附錄 A 各回合計時與完整數值案例',level=1)
d.add_paragraph('以下仍只使用本人批次。計時單位為 ns/op；L1／T4 為一條相依鏈／四條獨立鏈，ld 表示 long double。每組 7 筆，原始秒數、checksum 與所有命令保留於 q1-timing-and-precision.txt。')
samples=defaultdict(list)
for line in (RUN/'q1-timing-and-precision.txt').read_text().splitlines():
    if line.startswith('sample,'):
        f=line.split(',');v=dict(x.split('=',1) for x in f[4:]);samples[tuple(f[1:4])].append((int(v['run']),float(v['ns_per_operation'])))
for op in ['add_sub','multiply','divide']:
    d.add_heading('Q1 '+op+' 的全部回合',level=2)
    rr=[]
    for mode in ['latency1','throughput4']:
        for typ in ['float','double','long_double']:
            vals=sorted(samples[(op,mode,typ)]);assert len(vals)==7
            rr.append([typ.replace('long_double','ld'),'L1' if mode=='latency1' else 'T4']+[f'{v:.6f}' for _,v in vals])
    table(['型態','模式']+[str(i) for i in range(1,8)],rr,widths=[.55,.4]+[.79]*7,size=8.5)
d.add_heading('Q2 default 與 libm 的完整案例',level=2)
d.add_paragraph('兩版本共 60 次呼叫；30 個型態／輸入組合的 sqrt、errno、invalid、inexact、signbit、isnan 逐項一致，故合併顯示。特殊值類別與原始輸入／hexfloat 在 log 中完整保存。T／F 表示 true／false。')
for typ in ['float','double','long_double']:
    d.add_heading(typ.replace('_',' ')+' 的 sqrt',level=3)
    rr=[]
    for a in q2a:
        if a['type']==typ:rr.append([a['case'],a['sqrt'],a['errno']]+['T' if a[k]=='true' else 'F' for k in ['FE_INVALID','FE_INEXACT','negative_sign','isnan']])
    table(['輸入類別','sqrt 結果','errno','IV','IX','sign','NaN'],rr,widths=[1.3,2.7,.5,.5,.5,.5,.5],size=9)
d.add_heading('Q4 四模式與相鄰值',level=2)
q4=(RUN/'q4-formatting.txt').read_text().splitlines();mode=None;case=None;rr=[];a={}
for line in q4:
    if line.startswith('mode='):mode=line.split(',')[0].split('=',1)[1]
    elif line.startswith('case='):case=line.split('=',1)[1]
    elif line.startswith('printf_9=') and mode:
        a=dict(x.split('=',1) for x in line.split(','));rr.append([mode,case,a['printf_9'],a['iostream_9']])
assert len(rr)==16
table(['模式','案例','printf 9 位','iostream 9 位'],rr,widths=[1,2.3,1.6,1.6],size=9.5)
d.add_heading('附錄 B 證據索引與重現資料',level=1)
d.add_paragraph('本人結果目錄：results/linux-20261004-201349/。Windows snapshot：results/q0-windows/20261004-200712/。原始檔案與 screenshot 一起提交版本庫；截圖僅由 Word 原生裁去空白或顯示局部，未重繪數字。未重複貼入的圖片仍保存在 screenshot。')
d.add_heading('Q5 與共用來源雜湊',level=2)
for line in q5hash.splitlines():
    digest,name=line.split(maxsplit=1)
    x=d.add_paragraph(name+'\n'+digest)
    for r in x.runs:r.font.size=Pt(9)
d.add_paragraph('Q5 binary SHA-256：'+(RUN/'q5-binary-sha256.txt').read_text().split()[0])
d.add_heading('已插入的圖與原檔',level=2)
table(['圖／證據','原始截圖檔名'],[[f['caption'].split(' ')[0]+' '+f['caption'].split(' ')[1],f['file']] for f in figures],widths=[1.2,5.3],size=9)
d.add_paragraph('其餘 screenshot：202348 為 Q2 路徑標記；202909／202959 是 Q4 更長的重複顯示，關鍵數值已在圖 4-1 與完整案例表呈現。所有 21 張原圖均保留，避免以重複畫面增加篇幅。')

# Existing paragraph/page style structure is preserved; figures never split from captions.
for st in ['Title','Heading 1','Heading 2','Heading 3']:
    d.styles[st].font.color.rgb=RGBColor(0,0,0)
for x in d.paragraphs:
    if x.text.startswith('【待填】') or '【截圖位置】' in x.text:raise AssertionError(x.text)
d.save(DOC)
manifest={'student_run':RUN.name,'windows_snapshot':'20261004-200712','figures':figures,'raw_q1_samples':126,'accuracy_cases':12,'q2_calls':60,'q4_mode_cases':16,'mac_student_q5':'pending'}
(ROOT/'docs/student_report_evidence.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('Filled report with student data and',len(figures),'actual screenshots.')
